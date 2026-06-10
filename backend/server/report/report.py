from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from ._report import ReportData, ReportRenderer, ReportRepository, ReportService
from models.internal import (
    Attendance,
    Event,
    EventTagLink,
    Notification,
    Registration,
    Tag,
    User,
)


class SqlModelReportRepository(ReportRepository):
    def __init__(self, session: Session) -> None:
        self._session = session

    def collect(self, date_from: date, date_to: date) -> ReportData:
        dt_from = datetime.combine(date_from, datetime.min.time())
        dt_to = datetime.combine(date_to, datetime.max.time())

        users = self._session.exec(
            select(User).where(User.created_at >= dt_from, User.created_at <= dt_to)
        ).all()
        events = self._session.exec(
            select(Event).where(Event.created_at >= dt_from, Event.created_at <= dt_to)
        ).all()
        tag_links = self._session.exec(
            select(EventTagLink, Tag, Event)
            .join(Tag, EventTagLink.tag_id == Tag.id)
            .join(Event, EventTagLink.event_id == Event.id)
            .where(
                EventTagLink.event_id.in_(
                    select(Event.id).where(
                        Event.created_at >= dt_from,
                        Event.created_at <= dt_to,
                    )
                )
            )
        ).all()
        registrations = self._session.exec(
            select(Registration, User, Event)
            .join(User, Registration.user_id == User.id)
            .join(Event, Registration.event_id == Event.id)
            .where(
                Registration.created_at >= dt_from,
                Registration.created_at <= dt_to,
            )
        ).all()
        attendances = self._session.exec(
            select(Attendance, User, Event)
            .join(User, Attendance.user_id == User.id)
            .join(Event, Attendance.event_id == Event.id)
            .where(
                Attendance.created_at >= dt_from,
                Attendance.created_at <= dt_to,
            )
        ).all()
        notifications = self._session.exec(
            select(Notification).where(
                Notification.created_at >= dt_from,
                Notification.created_at <= dt_to,
            )
        ).all()

        return ReportData(
            users=list(users),
            events=list(events),
            tag_links=list(tag_links),
            registrations=list(registrations),
            attendances=list(attendances),
            notifications=list(notifications),
        )


class ExcelReportRenderer(ReportRenderer):
    def render(self, data: ReportData) -> BytesIO:
        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4472C4",
            end_color="4472C4",
            fill_type="solid",
        )
        header_alignment = Alignment(horizontal="center", vertical="center")

        def style_header(ws, headers):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

        def auto_width(ws):
            for col in ws.columns:
                max_length = 0
                column_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

        ws_users = wb.active
        ws_users.title = "Пользователи"
        style_header(
            ws_users,
            [
                "ID",
                "Никнейм",
                "Имя",
                "Фамилия",
                "Телефон",
                "Роль",
                "Баллы",
                "Компания",
                "Создан",
            ],
        )
        for row_num, user in enumerate(data.users, 2):
            ws_users.cell(row=row_num, column=1, value=user.id)
            ws_users.cell(row=row_num, column=2, value=user.nickname)
            ws_users.cell(row=row_num, column=3, value=user.firstname)
            ws_users.cell(row=row_num, column=4, value=user.lastname)
            ws_users.cell(row=row_num, column=5, value=user.phone)
            ws_users.cell(row=row_num, column=6, value=user.role.value)
            ws_users.cell(row=row_num, column=7, value=user.points)
            ws_users.cell(row=row_num, column=8, value=user.company)
            ws_users.cell(
                row=row_num,
                column=9,
                value=user.created_at.strftime("%Y-%m-%d %H:%M")
                if user.created_at
                else "",
            )
        auto_width(ws_users)

        ws_events = wb.create_sheet("События")
        style_header(
            ws_events,
            ["ID", "Название", "Описание", "Дата", "Баллы", "Ссылка", "Архив", "Создано"],
        )
        for row_num, event in enumerate(data.events, 2):
            ws_events.cell(row=row_num, column=1, value=event.id)
            ws_events.cell(row=row_num, column=2, value=event.title)
            ws_events.cell(row=row_num, column=3, value=event.description)
            ws_events.cell(
                row=row_num,
                column=4,
                value=event.date.strftime("%Y-%m-%d %H:%M") if event.date else "",
            )
            ws_events.cell(row=row_num, column=5, value=event.points)
            ws_events.cell(row=row_num, column=6, value=event.link)
            ws_events.cell(row=row_num, column=7, value="Да" if event.is_archived else "Нет")
            ws_events.cell(
                row=row_num,
                column=8,
                value=event.created_at.strftime("%Y-%m-%d %H:%M")
                if event.created_at
                else "",
            )
        auto_width(ws_events)

        ws_tags = wb.create_sheet("Теги")
        style_header(ws_tags, ["ID тега", "Название тега", "ID события", "Название события"])
        for row_num, (_link, tag, event) in enumerate(data.tag_links, 2):
            ws_tags.cell(row=row_num, column=1, value=tag.id)
            ws_tags.cell(row=row_num, column=2, value=tag.title)
            ws_tags.cell(row=row_num, column=3, value=event.id)
            ws_tags.cell(row=row_num, column=4, value=event.title)
        auto_width(ws_tags)

        ws_regs = wb.create_sheet("Регистрации")
        style_header(
            ws_regs,
            ["ID пользователя", "Никнейм", "ID события", "Название события", "Дата регистрации"],
        )
        for row_num, (reg, user, event) in enumerate(data.registrations, 2):
            ws_regs.cell(row=row_num, column=1, value=user.id)
            ws_regs.cell(row=row_num, column=2, value=user.nickname)
            ws_regs.cell(row=row_num, column=3, value=event.id)
            ws_regs.cell(row=row_num, column=4, value=event.title)
            ws_regs.cell(
                row=row_num,
                column=5,
                value=reg.created_at.strftime("%Y-%m-%d %H:%M") if reg.created_at else "",
            )
        auto_width(ws_regs)

        ws_att = wb.create_sheet("Посетители")
        style_header(
            ws_att,
            ["ID пользователя", "Никнейм", "ID события", "Название события", "Дата отметки"],
        )
        for row_num, (att, user, event) in enumerate(data.attendances, 2):
            ws_att.cell(row=row_num, column=1, value=user.id)
            ws_att.cell(row=row_num, column=2, value=user.nickname)
            ws_att.cell(row=row_num, column=3, value=event.id)
            ws_att.cell(row=row_num, column=4, value=event.title)
            ws_att.cell(
                row=row_num,
                column=5,
                value=att.created_at.strftime("%Y-%m-%d %H:%M") if att.created_at else "",
            )
        auto_width(ws_att)

        ws_notif = wb.create_sheet("Уведомления")
        style_header(ws_notif, ["ID", "Заголовок", "Текст", "Создано"])
        for row_num, notification in enumerate(data.notifications, 2):
            ws_notif.cell(row=row_num, column=1, value=notification.id)
            ws_notif.cell(row=row_num, column=2, value=notification.title)
            ws_notif.cell(row=row_num, column=3, value=notification.body)
            ws_notif.cell(
                row=row_num,
                column=4,
                value=notification.created_at.strftime("%Y-%m-%d %H:%M")
                if notification.created_at
                else "",
            )
        auto_width(ws_notif)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output


class DefaultReportService(ReportService):
    def __init__(self, repository: ReportRepository, renderer: ReportRenderer) -> None:
        self._repository = repository
        self._renderer = renderer

    def generate(self, date_from: date, date_to: date) -> BytesIO:
        return self._renderer.render(self._repository.collect(date_from, date_to))
