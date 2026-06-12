from datetime import date

from openpyxl import load_workbook

from server.report.report import (
    DefaultReportService,
    ExcelReportRenderer,
    SqlModelReportRepository,
)


def test_report_generation_includes_company_name(db_session, user_factory):
    user_factory(company="Report Company")

    report = DefaultReportService(
        repository=SqlModelReportRepository(db_session),
        renderer=ExcelReportRenderer(),
    ).generate(date(2000, 1, 1), date(2100, 1, 1))

    workbook = load_workbook(report)
    users_sheet = workbook["Пользователи"]

    assert users_sheet.cell(row=2, column=8).value == "Report Company"
